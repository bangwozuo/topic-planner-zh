# -*- coding: utf-8 -*-
"""
assettools —— bangwozuo 数字员工资产的共享工具库（单文件，零业务逻辑）。

各技能 scripts/ 通过 sys.path 引入本文件，用于：
  - 依赖检查（缺失时打印修复命令，不静默失败）
  - 数据读写（CSV / Excel，含条件格式）
  - 统计判定（环比 / 同比 / Z-score / IQR / EWMA / 线性趋势）
  - 打分（多因子加权）
  - 出图（折线 / 柱状 / 饼图，中文字体已配好）
  - 出文档（Word）

设计原则：
  1. 只提供**通用能力**，业务规则一律写在调用方脚本里
  2. 所有输出落到真实文件，返回绝对路径
  3. 中文可正常渲染（matplotlib 字体 + openpyxl 编码）
"""
from __future__ import annotations

import csv
import json
import math
import os
import subprocess
import sys
from datetime import datetime

# ---------------------------------------------------------------- 依赖

_REQUIRED = {
    "pandas": "pandas",
    "openpyxl": "openpyxl",
    "matplotlib": "matplotlib",
    "docx": "python-docx",
}


def need(*keys, install_hint: bool = True):
    """检查依赖；缺失则打印修复命令并退出（退出码 2），绝不静默继续。"""
    missing = []
    for k in keys:
        try:
            __import__(k)
        except ImportError:
            missing.append(_REQUIRED.get(k, k))
    if missing:
        pkgs = " ".join(sorted(set(missing)))
        print(f"[依赖缺失] 需要: {pkgs}", file=sys.stderr)
        if install_hint:
            exe = sys.executable or "python3"
            print(f"[修复命令] \"{exe}\" -m pip install {pkgs}", file=sys.stderr)
        sys.exit(2)


# ---------------------------------------------------------------- IO


def read_table(path):
    """读 CSV / XLSX → list[dict]。自动识别编码与扩展名。"""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        need("openpyxl")
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        head = [str(h) if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
        return [dict(zip(head, r)) for r in rows[1:]]
    for enc in ("utf-8-sig", "utf-8", "gbk"):
        try:
            with open(path, encoding=enc, newline="") as f:
                return [dict(r) for r in csv.DictReader(f)]
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法解码文件: {path}")


def write_json(obj, path):
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=2)
    return os.path.abspath(path)


def read_json(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def write_excel(path, sheets, highlights=None, widths=None):
    """写多 sheet Excel。

    sheets:     {"Sheet 名": [ {列:值}, ... ], ...}
    highlights: {"Sheet 名": {"列名": "条件"}}
                条件支持 '>N' '<N' '>=N' '<=N' 'contains:xxx'
    widths:     {"Sheet 名": {"列名": 20}}
    返回产物绝对路径。
    """
    need("openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F5597")
    warn_fill = PatternFill("solid", fgColor="FFC7CE")
    ok_fill = PatternFill("solid", fgColor="C6EFCE")

    for sname, rows in sheets.items():
        ws = wb.create_sheet(title=sname[:31])
        if not rows:
            ws.append(["(无数据)"])
            continue
        cols = list(rows[0].keys())
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for r in rows:
            ws.append([r.get(c, "") for c in cols])

        rule = (highlights or {}).get(sname, {})
        for ri, r in enumerate(rows, start=2):
            for ci, c in enumerate(cols, start=1):
                if c in rule and _hit(r.get(c), rule[c]):
                    ws.cell(row=ri, column=ci).fill = warn_fill
                    ws.cell(row=ri, column=ci).font = Font(color="9C0006")

        w = (widths or {}).get(sname, {})
        for ci, c in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(ci)].width = w.get(c, _auto_w(c, rows))
        ws.freeze_panes = "A2"
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb.save(path)
    return os.path.abspath(path)


def _auto_w(col, rows):
    """列宽自适应：中文按 2 个字符宽计。"""
    def w(s):
        s = str(s)
        return sum(2 if ord(ch) > 0x2E80 else 1 for ch in s)
    m = w(col)
    for r in rows[:200]:
        m = max(m, w(r.get(col, "")))
    return min(max(m + 2, 8), 48)


def _hit(val, cond):
    try:
        if cond.startswith("contains:"):
            return cond.split(":", 1)[1] in str(val)
        for op in (">=", "<=", ">", "<", "=="):
            if cond.startswith(op):
                num = float(cond[len(op):])
                return eval(f"float(val) {op} num")  # noqa: S307 —— 条件来自本仓脚本
    except (TypeError, ValueError):
        return False
    return False


# ---------------------------------------------------------------- 统计判定


def pct_change(cur, prev):
    """环比。prev 为 0 时返回 None（不制造无穷大）。"""
    if prev in (0, None) or cur is None:
        return None
    return (cur - prev) / abs(prev) * 100.0


def yoy(cur, same_period_last_year):
    """同比。"""
    return pct_change(cur, same_period_last_year)


def mean(xs):
    xs = [x for x in xs if x is not None]
    return sum(xs) / len(xs) if xs else None


def stdev(xs):
    """样本标准差（n-1）。"""
    xs = [x for x in xs if x is not None]
    if len(xs) < 2:
        return None
    m = mean(xs)
    return math.sqrt(sum((x - m) ** 2 for x in xs) / (len(xs) - 1))


def zscore(x, xs):
    """Z-score。样本 < 14 时应改用 IQR（标准差在小样本上不稳定）。"""
    s = stdev(xs)
    m = mean(xs)
    if not s:
        return None
    return (x - m) / s


def quantile(xs, q):
    """线性插值分位数，q ∈ [0,1]。"""
    xs = sorted(x for x in xs if x is not None)
    if not xs:
        return None
    if len(xs) == 1:
        return xs[0]
    pos = q * (len(xs) - 1)
    lo = int(math.floor(pos))
    hi = int(math.ceil(pos))
    if lo == hi:
        return xs[lo]
    return xs[lo] + (xs[hi] - xs[lo]) * (pos - lo)


def iqr_outliers(xs, k=1.5):
    """IQR 离群检测（小样本友好）。返回 (下界, 上界)。"""
    q1, q3 = quantile(xs, 0.25), quantile(xs, 0.75)
    if q1 is None or q3 is None:
        return None, None
    iqr = q3 - q1
    return q1 - k * iqr, q3 + k * iqr


def ewma(xs, alpha=0.3):
    """指数加权移动平均，用于趋势平滑。"""
    xs = [x for x in xs if x is not None]
    if not xs:
        return []
    out = [xs[0]]
    for x in xs[1:]:
        out.append(alpha * x + (1 - alpha) * out[-1])
    return out


def trend_slope(ys):
    """最小二乘斜率（每期变化量）。用于判断趋势方向。"""
    ys = [y for y in ys if y is not None]
    n = len(ys)
    if n < 2:
        return None
    xs = list(range(n))
    mx, my = mean(xs), mean(ys)
    den = sum((x - mx) ** 2 for x in xs)
    if not den:
        return None
    return sum((xs[i] - mx) * (ys[i] - my) for i in range(n)) / den


def consecutive_runs(flags):
    """统计连续 True 的最长长度与所有 (start, end) 区间。"""
    best, runs, start = 0, [], None
    for i, f in enumerate(flags):
        if f and start is None:
            start = i
        elif not f and start is not None:
            runs.append((start, i - 1))
            best = max(best, i - start)
            start = None
    if start is not None:
        runs.append((start, len(flags) - 1))
        best = max(best, len(flags) - start)
    return best, runs


# ---------------------------------------------------------------- 打分


def minmax_score(value, lo, hi, higher_is_better=True):
    """把原始值线性映射到 0-100。超出区间按端点截断。"""
    if value is None or hi == lo:
        return 0.0
    r = (value - lo) / (hi - lo)
    r = max(0.0, min(1.0, r))
    return round((r if higher_is_better else 1 - r) * 100, 1)


def weighted_score(parts, weights):
    """多因子加权打分。parts/weights 为 {因子: 值/权重}。返回 (总分, 明细)。"""
    tot_w = sum(weights.values()) or 1
    detail, total = {}, 0.0
    for k, v in parts.items():
        w = weights.get(k, 0)
        c = v * w
        detail[k] = {"得分": round(v, 1), "权重": w, "加权": round(c, 1)}
        total += c
    return round(total / tot_w, 1), detail


def quintile_rank(value, all_values):
    """五分位分层（1-5，5 为最高）。用于 RFM 类分层。"""
    qs = [quantile(all_values, i / 5) for i in range(1, 5)]
    rank = 1
    for q in qs:
        if q is not None and value > q:
            rank += 1
    return min(rank, 5)


# ---------------------------------------------------------------- 出图

_MPL_READY = False


def _mpl():
    """初始化 matplotlib（中文字体 + 后台渲染）。"""
    global _MPL_READY
    need("matplotlib")
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    if not _MPL_READY:
        plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
        plt.rcParams["axes.unicode_minus"] = False
        plt.rcParams["figure.dpi"] = 110
        _MPL_READY = True
    return plt


PALETTE = ["#2F5597", "#C00000", "#548235", "#BF8F00", "#7030A0", "#00838F"]


def line_chart(path, x, series, title="", xlabel="", ylabel="", y2=None,
               y2label="", figsize=(9, 4.6)):
    """折线图。series: {标签: [值]}；y2: 双轴右侧的 {标签: [值]}。"""
    plt = _mpl()
    fig, ax = plt.subplots(figsize=figsize)
    for i, (label, ys) in enumerate(series.items()):
        ax.plot(x, ys, marker="o", ms=3, lw=1.8, color=PALETTE[i % len(PALETTE)], label=label)
    ax.set_title(title, fontsize=13, fontweight="bold")
    ax.set_xlabel(xlabel, fontsize=10)
    ax.set_ylabel(ylabel, fontsize=10)
    ax.grid(alpha=0.25, ls="--")
    lines, labels = ax.get_legend_handles_labels()
    if y2:
        ax2 = ax.twinx()
        j = len(series)
        for k, (label, ys) in enumerate(y2.items()):
            ax2.plot(x, ys, marker="s", ms=3, lw=1.6, ls="--",
                     color=PALETTE[(j + k) % len(PALETTE)], label=label)
        ax2.set_ylabel(y2label or "", fontsize=10)
        l2, lb2 = ax2.get_legend_handles_labels()
        lines, labels = lines + l2, labels + lb2
    ax.legend(lines, labels, loc="upper center", bbox_to_anchor=(0.5, -0.18),
              ncol=min(len(labels), 4), fontsize=9, frameon=False)
    fig.tight_layout()
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return os.path.abspath(path)


def bar_chart(path, labels, values, title="", xlabel="", ylabel="", horizontal=False,
              figsize=(9, 4.6)):
    """柱状图。"""
    plt = _mpl()
    fig, ax = plt.subplots(figsize=figsize)
    if horizontal:
        ax.barh(labels, values, color=PALETTE[0], alpha=0.85)
        ax.set_xlabel(ylabel, fontsize=10)
    else:
        ax.bar(labels, values, color=PALETTE[0], alpha=0.85)
        ax.set_ylabel(ylabel, fontsize=10)
    ax.set_title(title, fontsize=13, fontweight="bold")
    if xlabel and not horizontal:
        ax.set_xlabel(xlabel, fontsize=10)
    ax.grid(alpha=0.25, ls="--", axis="y" if not horizontal else "x")
    plt.setp(ax.get_xticklabels(), rotation=20 if not horizontal else 0, ha="right" if not horizontal else "left")
    fig.tight_layout()
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return os.path.abspath(path)


def pie_chart(path, labels, values, title="", figsize=(6.4, 5)):
    """饼图（占比结构）。"""
    plt = _mpl()
    fig, ax = plt.subplots(figsize=figsize)
    ax.pie(values, labels=labels, autopct="%1.1f%%", startangle=90,
           colors=PALETTE[:len(labels)], textprops={"fontsize": 9})
    ax.set_title(title, fontsize=13, fontweight="bold")
    fig.tight_layout()
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return os.path.abspath(path)


# ---------------------------------------------------------------- 出文档


def write_docx(path, title, sections, subtitle=""):
    """写 Word 报告。

    sections: [ {"heading": "一、xxx", "paras": ["..."], "table": {"cols":[...], "rows":[[...]]},
                 "image": "/abs/x.png", "bullets": ["..."]} , ... ]
    返回产物绝对路径。
    """
    need("docx")
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor

    doc = Document()
    doc.add_heading(title, level=0)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.color.rgb = RGBColor(0x70, 0x70, 0x70)

    for sec in sections:
        if sec.get("heading"):
            doc.add_heading(sec["heading"], level=1)
        for t in sec.get("paras", []):
            doc.add_paragraph(t)
        for b in sec.get("bullets", []):
            doc.add_paragraph(b, style="List Bullet")
        tb = sec.get("table")
        if tb and tb.get("cols"):
            t = doc.add_table(rows=1, cols=len(tb["cols"]))
            t.style = "Light Grid Accent 1"
            for i, c in enumerate(tb["cols"]):
                cell = t.rows[0].cells[i]
                cell.text = str(c)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            for row in tb.get("rows", []):
                cells = t.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = "" if v is None else str(v)
        img = sec.get("image")
        if img and os.path.exists(img):
            doc.add_picture(img, width=Inches(5.9))

    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    doc.save(path)
    return os.path.abspath(path)


# ---------------------------------------------------------------- CLI 辅助


def emit(result):
    """统一返回 JSON（供调用方/智能体读取产物路径）。"""
    print(json.dumps(result, ensure_ascii=False, indent=2))


def ensure_outdir(path):
    os.makedirs(path, exist_ok=True)
    return os.path.abspath(path)


def stamp():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")
